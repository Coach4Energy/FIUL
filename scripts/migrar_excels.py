"""
FIUL · Etapa 1 · Migración de los 3 Excel a la tabla `alimentos`.

Lee Molifit, FitProtein y Suplementos, normaliza los valores (números como
texto, fechas corruptas, precios en 0) y arma una planilla de revisión para
que Juanba confirme la fuente de cada fila antes de cargarla a Supabase.

Uso:  python migrar_excels.py <molifit.xlsx> <fitprotein.xlsx> <suplementos.xlsx> <salida.xlsx>
"""
import sys, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

MICROS = ["sodio_mg", "calcio_mg", "hierro_mg", "vit_a_ug", "vit_c_mg", "vit_b6_mg", "vit_b12_ug"]

# Marcas → la fuente propuesta es el rótulo del producto
MARCAS = ["APIS", "Fargo", "Casancrem", "Campagnola", "Nescafé", "Serenísima", "Archimboldo",
          "BeePure", "LifeExtension", "Redoxon", "Nutremax", "FuXion", "STAR", "PÔNT", "Pizarro"]

# Grupo AIS propuesto (marco del Australian Institute of Sport)
AIS = {
    "Creatina monohidrato": ("A", "Suplemento de rendimiento (creatina)"),
    "Omega 3 - LifeExtension": ("B", "AIS lista 'fish oils' en Grupo B, no en A"),
    "Vitamina D3+C - Redoxon": ("A", "Vitamina D = suplemento médico Grupo A (la vit. C sola sería Grupo B)"),
    "Nutremax Energy-Gel C/Cafe": ("A", "Gel deportivo + cafeína, ambos Grupo A"),
    "Nutremax Energy-Gel S/Cafe": ("A", "Gel deportivo, Grupo A"),
    "FuXion BioPro+ Sport": ("A", "Suplemento proteico, Grupo A"),
    "Platinum Whey | STAR NUTRITION": ("A", "Suplemento proteico, Grupo A"),
    "Barra de Proteínas (PÔNT, 70g)": ("A", "Barra deportiva / proteica, Grupo A"),
}


def num(v):
    """Convierte '12,2' → 12.2; None/'' → None."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return float(str(v).replace(",", ".").strip())


def leer_molifit(path):
    ws = openpyxl.load_workbook(path, data_only=True)["Almuerzos y Cenas"]
    seccion, filas = None, []
    mapa = {"PROTEINAS": "proteina", "GUARNICIONES": "guarnicion",
            "GRASAS Y ACEITES": "grasa", "DESAYUNO Y TE": "desayuno_te"}
    for r in ws.iter_rows(values_only=True):
        if r[0] in mapa:
            seccion = mapa[r[0]]; continue
        if r[0] in (None, "Alimento"):
            continue
        cat = seccion
        filas.append(dict(
            proveedor="Molifit", categoria=cat, nombre=r[0].strip(),
            porcion_g=num(r[1]), porcion_descripcion=None,
            porcion_fija=cat in ("proteina", "guarnicion"),
            estado="cocido" if cat in ("proteina", "guarnicion") else "listo_para_consumir",
            kcal=num(r[2]), proteina_g=num(r[3]), carbohidratos_g=num(r[4]), grasas_g=num(r[5]),
            fibra_g=num(r[6]), **dict(zip(MICROS, [num(x) for x in r[7:14]])),
            precio_original=r[14], grupo_ais=None, problemas=[], activo=True))
    return filas


def leer_fitprotein(path):
    ws = openpyxl.load_workbook(path, data_only=True)["Almuerzos y Cenas"]
    filas = []
    for r in ws.iter_rows(values_only=True):
        if r[0] in (None, "Alimento", "FitProtein"):
            continue
        prob, desc, pg = [], None, None
        p = r[1]
        if isinstance(p, datetime.datetime):
            if p.year == 2026 and p.month == 1 and p.day == 4:
                desc = "1/4"
                prob.append("Porción guardada como fecha (04/01/2026) por Excel; se interpretó como '1/4'. "
                            "Confirmar 1/4 de qué (¿torta?) y su peso en gramos.")
            else:
                desc = "6 (a confirmar)"
                prob.append(f"Porción guardada como fecha ({p:%d/%m/%Y}) = número 6 en Excel. "
                            "Confirmar si son 6 unidades y su peso en gramos (560 kcal parece la bolsa entera).")
        elif p is not None:
            desc = "1 unidad (a confirmar)"
            prob.append("Porción = 1, sin gramos. Confirmar peso.")
        if any(isinstance(x, str) for x in r[3:6]):
            prob.append("Macros guardados como texto con coma decimal (corregido).")
        prob.append("Sin fibra ni micronutrientes en la base original.")
        filas.append(dict(
            proveedor="FitProtein", categoria="postre" if r[0] not in ("Yogur Griego", "Chipa") else "snack",
            nombre=r[0].strip(), porcion_g=pg, porcion_descripcion=desc, porcion_fija=False,
            estado="listo_para_consumir",
            kcal=num(r[2]), proteina_g=num(r[3]), carbohidratos_g=num(r[4]), grasas_g=num(r[5]),
            fibra_g=num(r[6]), **dict(zip(MICROS, [num(x) for x in r[7:14]])),
            precio_original=r[14], grupo_ais=None, problemas=prob, activo=True))
    return filas


def leer_suplementos(path):
    ws = openpyxl.load_workbook(path, data_only=True)["SUPLEMENTOS"]
    filas = []
    for r in ws.iter_rows(values_only=True):
        if r[0] in (None, "Alimento", "Suplementos"):
            continue
        nombre = r[0].strip()
        es_sup = nombre in AIS
        prob, activo = [], True
        if es_sup:
            g, why = AIS[nombre]
            prob.append(f"Grupo AIS propuesto {g}: {why}.")
        else:
            prob.append("No es un suplemento según AIS (es un alimento); se carga como snack, sin grupo AIS.")
        if nombre.startswith("Platinum Whey"):
            activo = False
            prob.append("Reemplazado por FuXion BioPro+ Sport: se carga inactivo.")
        filas.append(dict(
            proveedor="Suplementos", categoria="suplemento" if es_sup else "snack",
            nombre=nombre, porcion_g=num(r[1]), porcion_descripcion=None, porcion_fija=False,
            estado="listo_para_consumir",
            kcal=num(r[2]), proteina_g=num(r[3]), carbohidratos_g=num(r[4]), grasas_g=num(r[5]),
            fibra_g=num(r[6]), **dict(zip(MICROS, [num(x) for x in r[7:14]])),
            precio_original=r[14], grupo_ais=AIS[nombre][0] if es_sup else None,
            problemas=prob, activo=activo))
    return filas


# Criterio de la base (decidido 12/10/2026): carbohidratos DISPONIBLES (sin fibra), como el rótulo argentino.
# Estas filas traían carbohidratos totales (la fibra incluida): sus kcal solo cierran restando la fibra.
CONVERTIR_A_DISPONIBLES = {"Brócoli", "Calabaza / Zapallo", "Remolacha", "Zanahoria Cubos", "Palta (Pulpa)",
                           "Fruta - Densidad Calórica Media (Manzana/Pera/Naranja/Kiwi)",
                           "Fruta - Densidad Calórica Alta, Otoño-Invierno (Mango/Banana/Higo)"}
# Las kcal cierran con los dos criterios: confirmar en la fuente si los carbohidratos incluyen la fibra.
VERIFICAR_CRITERIO = {"Lentejas (cocidas)", "Boniato"}


def completar(f):
    if f["nombre"] in CONVERTIR_A_DISPONIBLES:
        orig = f["carbohidratos_g"]
        f["carbohidratos_g"] = round(orig - (f["fibra_g"] or 0), 2)
        f["problemas"].append(f"Carbohidratos convertidos a disponibles (totales {orig:g} g − fibra {f['fibra_g']:g} g).")
    if f["nombre"] == "Vegetales Sous Vide":
        f["problemas"].append("En pausa: Juanba consulta qué verduras trae el mix y en qué proporción.")
    if f["nombre"] in VERIFICAR_CRITERIO:
        f["problemas"].append("Al validar, confirmar en la fuente si los carbohidratos incluyen la fibra; "
                              "si la incluyen, restarla.")
    # precio
    p = f.pop("precio_original")
    if p in (None, 0, 0.0):
        f["precio_ars"] = None
        f["problemas"].append("Sin precio (en la base estaba en 0 o vacío).")
    else:
        f["precio_ars"] = float(p)
    # valores todos en cero
    if not any([f["kcal"], f["proteina_g"], f["carbohidratos_g"], f["grasas_g"]]):
        if f["categoria"] != "suplemento":
            f["problemas"].append("Todos los macros en 0: falta cargar los valores reales antes de validar.")
    # fuente propuesta
    marca = any(m.lower() in f["nombre"].lower() for m in MARCAS)
    f["fuente_propuesta"] = "Rotulo" if marca or f["proveedor"] in ("FitProtein", "Suplementos") else "Nutrinfo"
    return f


def construir(filas, salida):
    wb = openpyxl.Workbook()
    arial = "Arial"
    H = Font(name=arial, bold=True, color="FFFFFF")
    HF = PatternFill("solid", fgColor="000000")
    AM = PatternFill("solid", fgColor="FFFF00")
    ROJO = PatternFill("solid", fgColor="F4CCCC")
    thin = Side(style="thin", color="BFBFBF")

    # --- Hoja de instrucciones
    ws0 = wb.active; ws0.title = "Cómo usar"
    texto = [
        ("REVISIÓN DE ALIMENTOS · FIUL · Etapa 1", True),
        (f"Generado el {datetime.date.today():%d/%m/%Y} a partir de Molifit, FitProtein y Suplementos.", False),
        ("", False),
        ("Qué tenés que completar (celdas amarillas de la hoja 'Alimentos'):", True),
        ("1. Fuente: elegí Nutrinfo, Rotulo u Open Food Facts. Viene una propuesta; cambiala si no es correcta.", False),
        ("2. Detalle de la fuente: dónde está el dato exacto. Es obligatorio en la base.", False),
        ("3. Validado: poné 'Sí' cuando los números de la fila estén verificados contra esa fuente.", False),
        ("", False),
        ("Ejemplo de fila validada:", True),
        ("Fuente = Rotulo · Detalle = 'Rótulo Pan Lactal Fargo, porción 36 g (2 rebanadas)' · Validado = Sí", False),
        ("Fuente = Nutrinfo · Detalle = 'Nutrinfo Vademécum: Pechuga de pollo cocida, sin piel' · Validado = Sí", False),
        ("", False),
        ("Columnas de control (no se editan):", True),
        ("kcal calculadas = 4·proteína + 4·carbohidratos + 9·grasas + 2·fibra. 'Diferencia' en rojo = más de 10% de desvío.", False),
        ("Problemas = lo que encontré al migrar cada fila. Las filas en rojo no se pueden validar hasta resolverlas.", False),
        ("", False),
        ("Solo las filas con Validado = Sí pasan a Supabase. El resto queda esperando.", True),
    ]
    for i, (t, b) in enumerate(texto, 1):
        c = ws0.cell(row=i, column=1, value=t); c.font = Font(name=arial, bold=b, size=12 if i == 1 else 10)
    ws0.column_dimensions["A"].width = 120

    # --- Hoja de alimentos
    ws = wb.create_sheet("Alimentos")
    cols = ["proveedor", "categoria", "nombre", "porcion_g", "porcion_descripcion", "porcion_fija", "estado",
            "kcal", "proteina_g", "carbohidratos_g", "grasas_g", "fibra_g", *MICROS, "precio_ars", "grupo_ais",
            "activo", "fuente", "fuente_detalle", "validado", "kcal_calculadas", "diferencia_pct", "problemas"]
    titulos = ["Proveedor", "Categoría", "Nombre", "Porción (g)", "Porción (descripción)", "Porción fija 200 g",
               "Estado", "kcal", "Proteína (g)", "Carbohidratos (g)", "Grasas (g)", "Fibra (g)", "Sodio (mg)",
               "Calcio (mg)", "Hierro (mg)", "Vit. A (µg)", "Vit. C (mg)", "Vit. B6 (mg)", "Vit. B12 (µg)",
               "Precio (ARS)", "Grupo AIS", "Activo", "Fuente", "Detalle de la fuente", "Validado",
               "kcal calculadas", "Diferencia (%)", "Problemas"]
    ws.append(titulos)
    for c in ws[1]:
        c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical="center")
    L = {k: openpyxl.utils.get_column_letter(i + 1) for i, k in enumerate(cols)}

    for f in filas:
        r = ws.max_row + 1
        for k in cols:
            if k in ("fuente", "fuente_detalle", "validado", "kcal_calculadas", "diferencia_pct", "problemas"):
                continue
            v = f[k]
            if isinstance(v, bool):
                v = "Sí" if v else "No"
            ws[f"{L[k]}{r}"] = v
        ws[f"{L['fuente']}{r}"] = f["fuente_propuesta"]
        ws[f"{L['fuente']}{r}"].comment = Comment("Propuesta automática: confirmar.", "Claude")
        ws[f"{L['kcal_calculadas']}{r}"] = (f"=ROUND(4*{L['proteina_g']}{r}+4*{L['carbohidratos_g']}{r}"
                                            f"+9*{L['grasas_g']}{r}+2*N({L['fibra_g']}{r}),1)")
        ws[f"{L['diferencia_pct']}{r}"] = (f"=IF({L['kcal']}{r}=0,\"\",ROUND(({L['kcal']}{r}"
                                           f"-{L['kcal_calculadas']}{r})/{L['kcal']}{r},3))")
        ws[f"{L['diferencia_pct']}{r}"].number_format = "0.0%"
        ws[f"{L['problemas']}{r}"] = " · ".join(f["problemas"])
        for k in ("fuente", "fuente_detalle", "validado"):
            ws[f"{L[k]}{r}"].fill = AM
        bloqueante = any(s.startswith(("Todos los macros", "Porción guardada", "Porción = 1")) for s in f["problemas"])
        if bloqueante:
            for k in ("nombre", "problemas"):
                ws[f"{L[k]}{r}"].fill = ROJO

    n = ws.max_row
    for row in ws.iter_rows(min_row=2, max_row=n):
        for c in row:
            c.font = Font(name=arial, size=10); c.border = Border(bottom=thin)
            c.alignment = Alignment(vertical="top", wrap_text=c.column_letter == L["problemas"])
        ws[f"{L['precio_ars']}{row[0].row}"].number_format = '$#,##0;-$#,##0;"-"'

    # desvío > 10% en rojo
    from openpyxl.formatting.rule import FormulaRule
    d = L["diferencia_pct"]
    ws.conditional_formatting.add(f"{d}2:{d}{n}",
        FormulaRule(formula=[f'AND({d}2<>"",ABS({d}2)>0.1)'], fill=ROJO, font=Font(name=arial, color="9C0006", bold=True)))

    dv1 = DataValidation(type="list", formula1='"Nutrinfo,Rotulo,Open Food Facts"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True)
    ws.add_data_validation(dv1); ws.add_data_validation(dv2)
    dv1.add(f"{L['fuente']}2:{L['fuente']}{n}"); dv2.add(f"{L['validado']}2:{L['validado']}{n}")

    anchos = {"nombre": 42, "problemas": 70, "fuente_detalle": 40, "porcion_descripcion": 16, "estado": 18,
              "categoria": 13, "proveedor": 12}
    for k, col in L.items():
        ws.column_dimensions[col].width = anchos.get(k, 11)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{L['problemas']}{n}"

    # --- Resumen
    wr = wb.create_sheet("Resumen")
    wr.append(["Indicador", "Cantidad"])
    for c in wr[1]:
        c.font = H; c.fill = HF
    A = "Alimentos"
    resumen = [
        ("Filas migradas", f"=COUNTA('{A}'!C2:C{n})"),
        ("Validadas (Sí)", f"=COUNTIF('{A}'!{L['validado']}2:{L['validado']}{n},\"Sí\")"),
        ("Pendientes de validar", "=B2-B3"),
        ("Sin precio", f"=COUNTBLANK('{A}'!{L['precio_ars']}2:{L['precio_ars']}{n})"),
        ("kcal con desvío > 10%", f"=COUNTIF('{A}'!{d}2:{d}{n},\">0.1\")+COUNTIF('{A}'!{d}2:{d}{n},\"<-0.1\")"),
        ("Molifit", f"=COUNTIF('{A}'!A2:A{n},\"Molifit\")"),
        ("FitProtein", f"=COUNTIF('{A}'!A2:A{n},\"FitProtein\")"),
        ("Suplementos", f"=COUNTIF('{A}'!A2:A{n},\"Suplementos\")"),
    ]
    for t, fml in resumen:
        wr.append([t, fml])
    for row in wr.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=arial, size=10)
    wr.column_dimensions["A"].width = 28; wr.column_dimensions["B"].width = 12

    wb.active = 1
    wb.save(salida)


if __name__ == "__main__":
    mol, fit, sup, out = sys.argv[1:5]
    filas = [completar(f) for f in leer_molifit(mol) + leer_fitprotein(fit) + leer_suplementos(sup)]
    construir(filas, out)
    print(f"{len(filas)} filas → {out}")
